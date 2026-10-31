import os, sys, json, shutil, subprocess, traceback, zipfile
from playwright.sync_api import sync_playwright
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
S=os.path.join(ROOT,'tests','fixtures')
APP='file://'+os.path.join(ROOT,'build','GL-PDF.html')
OUT=os.path.join(ROOT,'build','test-out2'); shutil.rmtree(OUT,ignore_errors=True); os.makedirs(OUT)
results=[]
def ok(name, cond, info=''):
    results.append((name, bool(cond), info)); print(('PASS ' if cond else 'FAIL ')+name+(' — '+str(info)[:300] if info else ''), flush=True)
def run(cmd): return subprocess.run(cmd,capture_output=True,text=True).stdout
import re
def npages(p):
    m=re.search(r'Pages:\s+(\d+)', run(['pdfinfo',p])); return int(m.group(1)) if m else -1
def pyi(p, code):
    r=subprocess.run(['python3','-c','import pypdf,json,sys\nr=pypdf.PdfReader(sys.argv[1])\n'+code, p],capture_output=True,text=True); return (r.stdout+r.stderr).strip()

with sync_playwright() as pw:
    b=pw.chromium.launch()
    def newpage():
        ctx=b.new_context(viewport={'width':1366,'height':900},accept_downloads=True); pg=ctx.new_page(); pg._errs=[]
        pg.on('pageerror',lambda e:pg._errs.append(str(e)))
        pg.on('console',lambda m: m.type=='error' and pg._errs.append(m.text))
        pg.goto(APP); pg.wait_for_timeout(1200); return pg
    def openpdf(pg, f): pg.set_input_files('#fOpen', S+'/'+f); pg.wait_for_timeout(1600)
    def st(pg, js): return pg.evaluate('(()=>{const S=window.__glState();return '+js+'})()')
    def act(pg, a): pg.click('.tact[data-act='+a+']'); pg.wait_for_timeout(300)
    def dl(pg, fn, name, timeout=60000):
        with pg.expect_download(timeout=timeout) as d: fn()
        p=OUT+'/'+name; shutil.copy(d.value.path(), p); pg.wait_for_timeout(200); return p, d.value.suggested_filename
    def save(pg, name):
        pg.evaluate("document.querySelectorAll('.toast').forEach(t=>t.remove())"); pg.keyboard.press('Escape'); pg.wait_for_timeout(700)
        return dl(pg, lambda: pg.click('#bSave'), name)[0]
    def ok_btn(pg): pg.click('.modal .btn.pri')
    def fill(pg, name, val):
        loc=pg.locator('.modal [name="'+name+'"]')
        t=loc.evaluate('e=>e.tagName+":"+(e.type||"")')
        if t.startswith('SELECT'): loc.select_option(val)
        elif 'checkbox' in t: loc.set_checked(bool(val))
        else: loc.fill(str(val))
    def drag(pg, x, y, x2, y2): pg.mouse.move(x,y); pg.mouse.down(); pg.mouse.move(x2,y2,steps=6); pg.mouse.up(); pg.wait_for_timeout(300)
    def page_pt(pg, fx, fy, i=0):
        bb=pg.locator('.page').nth(i).bounding_box(); return bb['x']+bb['width']*fx, bb['y']+bb['height']*fy

    # ---------- panneau Outils ----------
    try:
        pg=newpage()
        ok('panneau Outils : 11 groupes', pg.locator('.tgroup').count()==11, pg.locator('.tgroup').count())
        openpdf(pg,'real.pdf')
        pg.click('.stab[data-panel=pages]'); pg.wait_for_timeout(500)
        ok('onglet Pages : miniatures', pg.locator('.th').count()==1)
        pg.click('.stab[data-panel=tools]')
        # en-tête / pied
        act(pg,'header'); fill(pg,'tl','<<fichier>>'); ok_btn(pg); pg.wait_for_timeout(300)
        ok('en-tête visible à l’écran', pg.locator('.deco div').count()>=2)
        act(pg,'watermark'); fill(pg,'text','BROUILLON'); ok_btn(pg); pg.wait_for_timeout(300)
        act(pg,'background'); fill(pg,'color','#fff3d6'); ok_btn(pg); pg.wait_for_timeout(300)
        act(pg,'props'); fill(pg,'title','Devis portail'); fill(pg,'author','G&L Créations'); ok_btn(pg); pg.wait_for_timeout(300)
        p=save(pg,'deco.pdf'); txt=run(['pdftotext','-layout',p,'-'])
        ok('pied « Page 1 sur 1 » dans le PDF', 'Page 1 sur 1' in txt, txt[-200:])
        ok('en-tête nom du fichier', 'real' in txt.split('\n')[0] or 'real' in txt)
        ok('filigrane dans le PDF', 'BROUILLON' in pyi(p,"print(r.pages[0].extract_text())"))
        info=run(['pdfinfo',p]); ok('propriétés (titre, auteur)', 'Devis portail' in info and 'G&L Cr' in info, info[:200])
        run(['pdftoppm','-r','30','-png','-singlefile',p,OUT+'/deco']); 
        from PIL import Image; c=Image.open(OUT+'/deco.png').convert('RGB').getpixel((5,300))
        ok('arrière-plan coloré', c!=(255,255,255) and c[2]<245, c)
        ok('aucune erreur (modifier)', not pg._errs, pg._errs[:2])
        # lien
        pg.click('.tact[data-act=link]'); x,y=page_pt(pg,.6,.12); drag(pg,x,y,x+120,y+30)
        fill(pg,'target','glcreations.fr'); ok_btn(pg); pg.wait_for_timeout(300)
        pg.keyboard.press('Escape')
        # recadrer
        act(pg,'crop'); x,y=page_pt(pg,.05,.05); drag(pg,x,y,*page_pt(pg,.95,.3)); ok_btn(pg); pg.wait_for_timeout(300)
        ok('recadrage visible', pg.locator('.cropmask').count()==4)
        p=save(pg,'link-crop.pdf')
        ok('lien enregistré', 'https://glcreations.fr' in pyi(p,"print([a.get_object().get('/A') for a in r.pages[0].get('/Annots',[])])"))
        cb=pyi(p,"print(r.pages[0].cropbox, r.pages[0].mediabox)"); ok('recadrage enregistré', cb.count('RectangleObject')==2 and cb.split(']')[0]!=cb.split(']')[1], cb)
        # numéro de page / page entière
        pg.click('#zPage'); pg.wait_for_timeout(300); ok('page entière', pg.inner_text('#zVal')!='')
        pg.context.close()
    except Exception as e: ok('bloc modifier', False, [l for l in traceback.format_exc().splitlines() if 'suite2.py' in l][-1:] + [traceback.format_exc().splitlines()[-1][:200]])

    # ---------- créer, combiner ----------
    try:
        pg=newpage()
        pg.click('#bCreate'); pg.click('.menu [data-c=blank]'); pg.wait_for_timeout(1500)
        ok('créer : page vierge', st(pg,'S.pages.length')==1 and st(pg,'S.name')=='Nouveau document.pdf')
        pg.click('#bCreate')
        with pg.expect_file_chooser() as fc: pg.click('.menu [data-c=images]')
        fc.value.set_files([S+'/scan.jpg', S+'/big.jpg']); pg.wait_for_timeout(400)
        ok_btn(pg); pg.wait_for_timeout(500)
        if pg.locator('.modal').count(): pg.locator('.modal .btn',has_text='Ne pas enregistrer').click()
        pg.wait_for_timeout(2500)
        ok('créer : PDF depuis 2 images', st(pg,'S.pages.length')==2, st(pg,'S.pages.length'))
        with pg.expect_file_chooser() as fc: act(pg,'combine')
        fc.value.set_files([S+'/real.pdf', S+'/multi.pdf', S+'/scan.jpg']); pg.wait_for_timeout(500)
        ok('combiner : liste de fichiers', pg.locator('.flist .fi').count()==3)
        pg.locator('.flist [data-dn="0"]').click(); pg.wait_for_timeout(100)
        ok_btn(pg); pg.wait_for_timeout(500)
        if pg.locator('.modal').count(): pg.locator('.modal .btn',has_text='Ne pas enregistrer').click()
        pg.wait_for_timeout(3000)
        n=st(pg,'S.pages.length'); ok('combiner : 2 + 1 + 3 + 1 pages', n==7, n)
        ok('aucune erreur (créer/combiner)', not pg._errs, pg._errs[:2])
        pg.context.close()
    except Exception as e: ok('bloc créer', False, [l for l in traceback.format_exc().splitlines() if 'suite2.py' in l][-1:] + [traceback.format_exc().splitlines()[-1][:200]])

    # ---------- organiser, extraire, diviser, insérer ----------
    try:
        pg=newpage(); openpdf(pg,'multi.pdf')
        act(pg,'organize'); pg.wait_for_timeout(800)
        ok('organiser : grille', pg.locator('.org .oc').count()==3)
        pg.locator('.org .oc').nth(0).click(); pg.locator('#oBar [data-o=right]').click(); pg.wait_for_timeout(400)
        pg.locator('.org .oc').nth(0).click(); pg.locator('#oBar [data-o=dup]').click(); pg.wait_for_timeout(500)
        ok('organiser : dupliquer', pg.locator('.org .oc').count()==4)
        pg.locator('.org .oc').nth(3).click(); pg.locator('#oBar [data-o=del]').click(); pg.wait_for_timeout(500)
        pg.locator('.org .oc').nth(1).click(); pg.locator('.org .oc').nth(2).click(modifiers=['Control'])
        p,_=dl(pg, lambda: pg.locator('#oBar [data-o=extract]').click(), 'org-extract.pdf')
        ok('organiser : extraire 2 pages', npages(p)==2, npages(p))
        pg.click('.modal [data-x]'); pg.wait_for_timeout(400)
        ok('organiser : rotations appliquées', st(pg,'S.pages[0].rot')==90 and st(pg,'S.pages.length')==3, (st(pg,'S.pages[0].rot'), st(pg,'S.pages.length')))
        pg.keyboard.press('Control+z'); pg.wait_for_timeout(300); ok('Ctrl+Z annule une action d’organisation', st(pg,'S.pages.length')==4)
        pg.keyboard.press('Control+y'); pg.wait_for_timeout(300)
        act(pg,'extract'); fill(pg,'r','1,3')
        p,_=dl(pg, lambda: ok_btn(pg), 'extract.pdf'); ok('extraire 1,3', npages(p)==2, npages(p))
        act(pg,'split'); fill(pg,'m','each')
        p,nm=dl(pg, lambda: ok_btn(pg), 'split.zip'); z=zipfile.ZipFile(p).namelist(); ok('diviser : 1 fichier par page (zip)', len(z)==3, z)
        act(pg,'insertfile'); fill(pg,'where','end')
        with pg.expect_file_chooser() as fc: ok_btn(pg)
        fc.value.set_files(S+'/real.pdf'); pg.wait_for_timeout(1500)
        ok('insérer depuis un fichier', st(pg,'S.pages.length')==4)
        act(pg,'rotateall'); ok('pivoter toutes les pages', st(pg,'S.pages[3].rot')==90)
        ok('aucune erreur (pages)', not pg._errs, pg._errs[:2])
        pg.context.close()
    except Exception as e: ok('bloc organiser', False, [l for l in traceback.format_exc().splitlines() if 'suite2.py' in l][-1:] + [traceback.format_exc().splitlines()[-1][:200]])

    # ---------- exporter ----------
    try:
        pg=newpage(); openpdf(pg,'real.pdf')
        pg.click('[data-t=check]'); pg.mouse.click(*page_pt(pg,.6,.05)); pg.keyboard.press('Escape')
        act(pg,'xword'); p,nm=dl(pg, lambda: ok_btn(pg), 'out.docx')
        import docx; d=docx.Document(p); t='\n'.join(x.text for x in d.paragraphs)
        ok('Word : fichier valide avec le texte', 'Devis n° 2026-118' in t and 'Société Martin' in t, t[:200])
        r=subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',OUT,p],capture_output=True,text=True,timeout=120); ok('Word : s’ouvre dans un traitement de texte', os.path.exists(OUT+'/out.pdf'), r.stderr[-200:])
        act(pg,'xword'); fill(pg,'mode','both'); p,_=dl(pg, lambda: ok_btn(pg), 'out2.docx'); ok('Word : mode texte + image', any(n.startswith('word/media/') for n in zipfile.ZipFile(p).namelist()))
        p,_=dl(pg, lambda: act(pg,'xexcel'), 'out.xlsx')
        import openpyxl; wb=openpyxl.load_workbook(p); ws=wb.worksheets[0]; vals=[c.value for row in ws.iter_rows() for c in row if c.value is not None]
        ok('Excel : montants en nombres', 2450 in vals and 890 in vals, vals[:12])
        p,_=dl(pg, lambda: act(pg,'xppt'), 'out.pptx', 90000)
        import pptx; pr=pptx.Presentation(p); ok('PowerPoint : 1 diapositive avec la page', len(pr.slides)==1 and len(pr.slides[0].shapes)>=1)
        act(pg,'ximg'); fill(pg,'fmt','png'); p,nm=dl(pg, lambda: ok_btn(pg), 'out.png'); from PIL import Image
        ok('Image PNG 150 ppp', Image.open(p).size[0] in range(1230,1250), Image.open(p).size)
        p,_=dl(pg, lambda: act(pg,'xrtf'), 'out.rtf'); rt=open(p,encoding='latin1').read(); ok('RTF', rt.startswith('{\\rtf1') and 'Devis' in rt)
        p,_=dl(pg, lambda: act(pg,'xtxt'), 'out.txt'); tx=open(p,encoding='utf-8-sig').read(); ok('Texte brut', 'Total TTC' in tx)
        ok('aucune erreur (export)', not pg._errs, pg._errs[:2])
        pg.context.close()
    except Exception as e: ok('bloc export', False, [l for l in traceback.format_exc().splitlines() if 'suite2.py' in l][-1:] + [traceback.format_exc().splitlines()[-1][:200]])

    # ---------- commentaires, paraphe, demander des signatures, formulaire ----------
    try:
        pg=newpage(); openpdf(pg,'real.pdf')
        pg.click('[data-t=note]'); pg.mouse.click(*page_pt(pg,.8,.1)); pg.keyboard.type('À vérifier'); pg.mouse.click(*page_pt(pg,.5,.2))
        pg.click('.stab[data-panel=comments]'); pg.wait_for_timeout(300)
        ok('liste des commentaires', pg.locator('.cmt-item').count()==1 and 'À vérifier' in pg.inner_text('#pComments'))
        pg.locator('.cmt-item').first.click(); pg.wait_for_timeout(200); ok('clic : sélection du commentaire', st(pg,'S.sel') is not None)
        pg.click('.stab[data-panel=tools]')
        act(pg,'initials'); pg.wait_for_timeout(200)
        ok('fenêtre paraphe', 'paraphe' in pg.inner_text('.modal h2'))
        pg.click('#tType'); pg.wait_for_timeout(500); pg.click('#sUse'); pg.wait_for_timeout(300); pg.mouse.click(*page_pt(pg,.9,.2)); pg.wait_for_timeout(200)
        im=st(pg,"S.pages[0].objs.filter(o=>o.type==='img')"); ok('paraphe posé (petit format)', len(im)==1 and im[0]['w']<=71, im[0]['w'] if im else None)
        act(pg,'reqsign'); fill(pg,'who','Client, Entreprise'); ok_btn(pg); pg.wait_for_timeout(600)
        ok('zones de signature placées', st(pg,"S.pages[0].objs.filter(o=>o.kind==='sig').length")==2)
        pg.click('.ctx [data-field=check]') if pg.locator('.ctx [data-field=check]').count() else None
        pg.click('.ctx [data-field=text]'); x,y=page_pt(pg,.1,.18); drag(pg,x,y,x+150,y+20)
        pg.fill('#cFName','nom_client'); pg.press('#cFName','Enter'); pg.wait_for_timeout(200)
        pg.keyboard.press('Escape'); pg.click('.tact[data-act=field]'); pg.click('.ctx [data-field=check]'); x,y=page_pt(pg,.1,.21); drag(pg,x,y,x+15,y+15)
        pg.keyboard.press('Escape'); pg.click('.tact[data-act=field]'); pg.click('.ctx [data-field=list]'); x,y=page_pt(pg,.4,.21); drag(pg,x,y,x+100,y+18)
        p=save(pg,'fields.pdf'); f=pyi(p,"print(json.dumps({k:str(v.get('/FT')) for k,v in (r.get_fields() or {}).items()}))")
        ok('champs créés (texte, case, liste, 2 signatures)', all(s in f for s in ['"nom_client": "/Tx"','/Btn','/Ch']) and f.count('/Sig')==2, f)
        pg2=newpage(); pg2.set_input_files('#fOpen',p); pg2.wait_for_timeout(1500)
        ok('PDF créé : champs remplissables dans G&L PDF', pg2.locator('.fl .ff').count()>=5 and pg2.locator('.fl .fsig').count()==2, pg2.locator('.fl .ff').count())
        pg2.context.close()
        ok('aucune erreur (signer/formulaire)', not pg._errs, pg._errs[:2])
        pg.context.close()
        pg=newpage(); openpdf(pg,'detect.pdf'); act(pg,'detect'); pg.wait_for_timeout(500)
        ok('détection automatique des champs', st(pg,"S.pages[0].objs.filter(o=>o.type==='field').length")==2, st(pg,"S.pages[0].objs.map(o=>o.name)"))
        # rechercher et biffer
        act(pg,'findredact'); fill(pg,'tel',True); fill(pg,'mail',True); ok_btn(pg); pg.wait_for_timeout(800)
        ok('rechercher et biffer (téléphone + e-mail)', st(pg,"S.pages[0].objs.filter(o=>o.type==='redact').length")==2)
        act(pg,'metaclean')
        p=save(pg,'redact.pdf'); t=run(['pdftotext',p,'-'])
        ok('biffage : texte supprimé du fichier', '06 20 81' not in t and 'contact@' not in t, t[:120])
        ok('biffage : page aplatie (aucun texte récupérable)', not t.strip(), t[:80])
        ok('informations cachées supprimées', 'G&L PDF' not in run(['pdfinfo',p]))
        pg.context.close()
    except Exception as e: ok('bloc signer', False, [l for l in traceback.format_exc().splitlines() if 'suite2.py' in l][-1:] + [traceback.format_exc().splitlines()[-1][:200]])

    # ---------- protéger, copie verrouillée, ouvrir un PDF protégé ----------
    try:
        pg=newpage(); openpdf(pg,'real.pdf')
        act(pg,'protect'); fill(pg,'user','gl2026'); fill(pg,'user2','gl2026'); fill(pg,'noEdit',True); fill(pg,'owner','patron'); ok_btn(pg); pg.wait_for_timeout(300)
        p=save(pg,'protect.pdf'); info=run(['pdfinfo','-upw','gl2026',p])
        ok('PDF protégé (AES-256, modification interdite)', 'AES-256' in info and 'change:no' in info, info[-160:])
        ok('sans mot de passe : illisible', 'Incorrect password' in subprocess.run(['pdfinfo',p],capture_output=True,text=True).stderr)
        pg2=newpage(); pg2.set_input_files('#fOpen',p); pg2.wait_for_timeout(1200)
        ok('ouverture : mot de passe demandé', pg2.locator('.modal input[name=pw]').count()==1)
        pg2.fill('.modal input[name=pw]','faux'); pg2.click('.modal .btn.pri'); pg2.wait_for_timeout(1200)
        ok('mauvais mot de passe refusé', pg2.locator('.modal input[name=pw]').count()==1)
        pg2.fill('.modal input[name=pw]','gl2026'); pg2.click('.modal .btn.pri'); pg2.wait_for_timeout(2000)
        ok('bon mot de passe : document ouvert', st(pg2,'S.pages.length')==1)
        p2=save(pg2,'unprot.pdf'); ok('réenregistré sans mot de passe', 'Encrypted:       no' in run(['pdfinfo',p2]))
        pg2.context.close()
        act(pg,'lockcopy'); fill(pg,'who','Vincent GUIRAUD'); p,_=dl(pg, lambda: ok_btn(pg), 'locked.pdf')
        info=run(['pdfinfo',p]); ok('copie verrouillée : lecture seule + journal', 'change:no' in info and npages(p)==2 and 'Journal du document' in run(['pdftotext',p,'-']), info[-150:])
        ok('aucune erreur (protéger)', not pg._errs, pg._errs[:2])
        pg.context.close()
    except Exception as e: ok('bloc protéger', False, [l for l in traceback.format_exc().splitlines() if 'suite2.py' in l][-1:] + [traceback.format_exc().splitlines()[-1][:200]])

    # ---------- compresser, OCR ----------
    try:
        pg=newpage(); openpdf(pg,'photo.pdf')
        act(pg,'compress'); ok_btn(pg); pg.wait_for_timeout(6000)
        msg=pg.inner_text('.modal') if pg.locator('.modal').count() else ''
        p,_=dl(pg, lambda: pg.locator('.modal .btn.pri').click(), 'compressed.pdf')
        ok('compression : fichier plus petit', os.path.getsize(p) < os.path.getsize(S+'/photo.pdf')*0.6, (os.path.getsize(S+'/photo.pdf'), os.path.getsize(p), msg[:100]))
        ok('compression : PDF valide', npages(p)==1)
        pg.context.close()
        pg=newpage(); openpdf(pg,'scan.pdf')
        act(pg,'ocr'); ok_btn(pg); pg.wait_for_timeout(1000)
        for _ in range(90):
            if not pg.locator('.busy').count(): break
            pg.wait_for_timeout(1000)
        n=st(pg,"(S.pages[0].ocr||[]).length"); ok('OCR : texte reconnu', n>=3, st(pg,"(S.pages[0].ocr||[]).map(l=>l.text)"))
        pg.keyboard.press('Control+f'); pg.keyboard.type('garde-corps'); pg.wait_for_timeout(1500)
        ok('OCR : texte cherchable', pg.inner_text('#sCount')=='1 sur 1', pg.inner_text('#sCount'))
        pg.keyboard.press('Escape')
        p=save(pg,'ocr.pdf'); ok('OCR : texte dans le PDF enregistré', 'Montant total' in run(['pdftotext',p,'-']), run(['pdftotext',p,'-'])[:200])
        ok('aucune erreur (compresser/OCR)', not pg._errs, pg._errs[:2])
        pg.context.close()
    except Exception as e: ok('bloc optimiser', False, [l for l in traceback.format_exc().splitlines() if 'suite2.py' in l][-1:] + [traceback.format_exc().splitlines()[-1][:200]])
    b.close()

fails=[r for r in results if not r[1]]
print('\n%d tests, %d réussis, %d échecs' % (len(results), len(results)-len(fails), len(fails)))
for f in fails: print(' ✗', f[0], '—', str(f[2])[:400])
sys.exit(1 if fails else 0)
